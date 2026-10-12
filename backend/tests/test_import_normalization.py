"""Generic import normalization: headers, blank rows, sentinels, years.

Regression fixture: ``ACMS 301.xlsx`` (read-only input; dry run only, its
data is never written). Expected invalid rows are DERIVED from the workbook
contents (department/brand actually missing) — never hardcoded row numbers,
and the importer itself contains no filename- or row-specific logic.

Every rule here must hold for any future CSV/XLS/XLSX with different column
names, ordering, capitalization, optional/extra columns and blank rows.
"""

import io
import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import Workbook, load_workbook  # noqa: E402

from app.routes import acs as acs_mod  # noqa: E402

REFERENCE_WORKBOOK = Path(__file__).resolve().parents[2] / "ACMS 301.xlsx"

# The 22 raw headers of the canonical reference workbook (header names only).
REFERENCE_HEADERS = [
    "college_name",
    "block_name",
    "floor_name",
    "department_name",
    "room_type",
    "room_number",
    "Brand / Company",
    "unit_number",
    "model",
    "serial_number",
    "capacity_ton",
    "type",
    "manufacturing_year",
    "installation_year",
    "star_rating",
    "power_factor",
    "status",
    "last_maintenance_type",
    "last_maintenance_date",
    "next_maintenance_date",
    "technician",
    "remarks",
]


@pytest.fixture()
def client(db_admin_client):
    return db_admin_client


class _Upload(io.BytesIO):
    """Minimal FileStorage stand-in for parser-level tests (no DB)."""

    def __init__(self, data, filename):
        super().__init__(data if isinstance(data, bytes) else data.encode("utf-8"))
        self.filename = filename


def _parse(body, filename="f.csv"):
    return acs_mod._read_import_rows(_Upload(body, filename))


def _payload(body, index=0, filename="f.csv"):
    _headers, rows = _parse(body, filename)
    assert rows, "expected at least one parsed row"
    return acs_mod._csv_payload(rows[index], index + 2)


def _csv(header, rows):
    lines = [",".join(header)] + [",".join("" if c is None else str(c) for c in row) for row in rows]
    return "\n".join(lines) + "\n"


def _workbook_bytes(header, rows):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(header)
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _post_dry(client, body, filename="f.csv"):
    data = body if isinstance(body, bytes) else body.encode("utf-8")
    return client.post(
        "/api/v1/assets/import?dry_run=true",
        data={"file": (io.BytesIO(data), filename)},
    )


def _report(response):
    payload = response.get_json()
    assert payload["success"] is True, payload
    return payload["data"]


# ── 4/5/6: make-column aliases ──────────────────────────────────────────────


@pytest.mark.parametrize("header", ["Brand / Company", "brand_company", "Manufacturer"])
def test_make_alias_headers_map_to_ac_make_company(header):
    values, _lookup = _payload(f"ac_id,{header}\nX-ALIAS,Voltas\n")
    assert values["ac_make_company"] == "Voltas"


# ── 7: header capitalization / mixed spacing ────────────────────────────────


def test_header_capitalization_and_spacing_are_normalized():
    body = "AC_ID,Brand / COMPANY,MODEL_Number,SERIAL\nX-CAPS,Voltas,TKM50,SN-CAPS\n"
    values, _lookup = _payload(body)
    assert values["ac_id"] == "X-CAPS"
    assert values["ac_make_company"] == "Voltas"
    assert values["ac_model_number"] == "TKM50"
    assert values["ac_serial_number"] == "SN-CAPS"


# ── 8: column ordering ──────────────────────────────────────────────────────


def test_column_order_does_not_matter():
    forward = _payload("ac_id,brand,model,capacity\nX-1,Daikin,FTKM50,1.5\n")
    shuffled = _payload("capacity,model,ac_id,brand\n1.5,FTKM50,X-1,Daikin\n")
    assert forward[0] == shuffled[0]
    assert forward[1] == shuffled[1]


# ── 9/10: extra and missing columns (dry run) ───────────────────────────────


def test_extra_columns_are_reported_not_mapped(client):
    body = _csv(["ac_id", "brand", "warranty_months"], [["X-EXTRA", "LG", 24]])
    data = _report(_post_dry(client, body))
    assert data["unknown_columns"] == ["warranty_months"]
    assert data["imported"] == 1 and data["invalid"] == 0


def test_missing_optional_columns_are_reported_not_fatal(client):
    body = _csv(["ac_id", "brand"], [["X-MISS", "Daikin"]])
    data = _report(_post_dry(client, body))
    for column in ("room_id", "power_rating_kw", "installation_date", "ac_model_number"):
        assert column in data["missing_fields"]
    # The make column must be recognized via the "brand" alias, not "missing".
    assert "ac_make_company" not in data["missing_fields"]
    assert data["imported"] == 1


# ── 11: missing required fields stay fatal ──────────────────────────────────


def test_missing_brand_is_invalid(client):
    body = _csv(["ac_id", "brand"], [["X-NOBRAND", ""]])
    data = _report(_post_dry(client, body))
    assert data["invalid"] == 1 and data["imported"] == 0
    assert "Brand/company is required" in data["errors"][0]["message"]


def test_missing_department_in_canonical_mode_is_invalid(client):
    body = _csv(
        ["college_name", "block_name", "department_name", "Brand / Company", "room_number"],
        [["NITTTRC", "L block", "", "VESTAR", "210"]],
    )
    data = _report(_post_dry(client, body))
    assert data["invalid"] == 1 and data["imported"] == 0
    assert "department_name is required" in data["errors"][0]["message"]


# ── 12/13: unknown/sentinel values normalize to NULL ────────────────────────


@pytest.mark.parametrize(
    "sentinel",
    ["NO DATA", "NA", "N/A", "n.a.", "Not Available", "UNKNOWN", "-"],
)
def test_sentinel_values_normalize_to_null(sentinel):
    values, _lookup = _payload(f"ac_id,brand,installation_year\nX-SNT,Voltas,{sentinel}\n")
    assert values.get("installation_year") is None


def test_misspelled_data_not_available_sentinel_normalizes_to_null():
    # Exact spelling found in the reference workbook (leading spaces + typo).
    body = "ac_id,brand,manufacturing_year\nX-TYPO,Voltas,     DATA NOT AVilabe\n"
    values, _lookup = _payload(body)
    assert values.get("manufacturing_year") is None


# ── 14: valid numeric years ─────────────────────────────────────────────────


@pytest.mark.parametrize("raw", ["2020", " 2020 ", "2020.0"])
def test_valid_year_strings_coerce_to_int(raw):
    values, _lookup = _payload(f"ac_id,brand,installation_year\nX-YR,Voltas,{raw}\n")
    assert values["installation_year"] == 2020


def test_excel_numeric_year_cell_coerces_to_int():
    body = _workbook_bytes(
        ["ac_id", "brand", "installation_year"],
        [["X-YR2", "Voltas", 2020]],
    )
    _headers, rows = _parse(body, "f.xlsx")
    values, _lookup = acs_mod._csv_payload(rows[0], 2)
    assert values["installation_year"] == 2020


# ── 15: invalid years stay invalid ──────────────────────────────────────────


@pytest.mark.parametrize("raw", ["20X9", "twenty", "2020.5", "99999"])
def test_invalid_years_are_rejected(raw):
    with pytest.raises(ValueError, match="invalid installation_year"):
        _payload(f"ac_id,brand,installation_year\nX-BAD,Voltas,{raw}\n")


# ── 2/3: blank rows (CSV and XLSX) ──────────────────────────────────────────


@pytest.mark.parametrize("filename", ["f.csv", "f.xlsx"])
def test_trailing_blank_rows_are_ignored(client, filename):
    header = ["ac_id", "brand", "model"]
    records = [["X-1", "Daikin", "M1"], ["X-2", "LG", "M2"]]
    blanks = [[None, None, None]] * 3
    body = (
        _workbook_bytes(header, records + blanks)
        if filename.endswith(".xlsx")
        else _csv(header, records + [["", "", ""]] * 3)
    )
    data = _report(_post_dry(client, body, filename))
    assert data["physical_rows"] == 5
    assert data["blank_rows_ignored"] == 3
    assert data["total_rows"] == 2
    assert data["invalid"] == 0 and data["imported"] == 2


def test_blank_rows_between_records_keep_original_row_numbers(client):
    header = ["ac_id", "brand"]
    rows = [
        ["X-1", "Daikin"],
        [None, None],
        [None, None],
        ["X-2", "LG"],
        ["X-3", "Sony"],
        [None, None],
    ]
    data = _report(_post_dry(client, _workbook_bytes(header, rows), "f.xlsx"))
    assert data["physical_rows"] == 6
    assert data["blank_rows_ignored"] == 3
    assert data["total_rows"] == 3
    assert data["imported"] == 3 and data["invalid"] == 0
    # Reported rows are the ORIGINAL spreadsheet rows (2, 5, 6), not indices.
    assert sorted(entry["row"] for entry in data["rows"]) == [2, 5, 6]


# ── 17: ambiguous duplicate alias columns ───────────────────────────────────


def test_duplicate_alias_columns_resolve_by_documented_precedence(client):
    # Three headers that all map to ac_make_company: precedence follows the
    # alias tuple order ("brand" before "company" before "brand_company"),
    # and none of them may be reported as unknown.
    body = "ac_id,Brand / Company,company,brand\nX-AMB,Fallback Co,Co Inc,Voltas\n"
    values, _lookup = _payload(body)
    assert values["ac_make_company"] == "Voltas"
    data = _report(_post_dry(client, body))
    assert data["unknown_columns"] == []
    assert data["recognized_columns"] == ["ac_id", "ac_make_company"]


# ── 16: CSV equivalent of the reference data ────────────────────────────────


def _reference_row(**overrides):
    base = {header: "" for header in REFERENCE_HEADERS}
    base.update(
        {
            "college_name": "NITTTRC",
            "block_name": "L block",
            "floor_name": "first",
            "department_name": "HOSTEL",
            "room_type": "Hostel Room",
            "room_number": "210",
            "Brand / Company": "VESTAR",
            "unit_number": "19",
            "model": "FTKM50",
            "capacity_ton": "1.5",
            "type": "Split",
            "manufacturing_year": "DATA NOT AVilabe",
            "installation_year": "NO DATA",
            "star_rating": "4",
            "status": "Working",
        }
    )
    base.update(overrides)
    return [base[header] for header in REFERENCE_HEADERS]


def test_csv_equivalent_data_gets_same_classification(client):
    blank = ["" for _ in REFERENCE_HEADERS]
    rows = [
        _reference_row(),  # Excel row 2 — valid, both years are sentinels
        _reference_row(unit_number="20", installation_year="2020"),  # row 3 — valid
        blank,  # row 4 — blank between records
        _reference_row(department_name=""),  # row 5 — invalid: no department
        _reference_row(**{"Brand / Company": ""}),  # row 6 — invalid: no brand
        blank,  # rows 7/8 — trailing blanks
        blank,
    ]
    body = _csv(REFERENCE_HEADERS, rows)

    # Parser level: sentinels became NULL, the numeric year coerced.
    values, _lookup = _payload(body, 0)
    assert values.get("manufacturing_year") is None
    assert values.get("installation_year") is None
    values, _lookup = _payload(body, 1)
    assert values["installation_year"] == 2020

    # Endpoint level: identical classification semantics to the workbook.
    data = _report(_post_dry(client, body))
    assert data["physical_rows"] == 7
    assert data["blank_rows_ignored"] == 3
    assert data["total_rows"] == 4
    assert data["imported"] == 2
    assert data["invalid"] == 2
    assert data["duplicates"] == 0
    assert data["unknown_columns"] == []
    messages = {entry["row"]: entry["message"] for entry in data["errors"]}
    assert messages == {
        5: "row 5: department_name is required",
        6: "row 6: Brand/company is required.",
    }


# ── 1: ACMS 301.xlsx regression (dry run, read-only) ────────────────────────


def _derive_invalid_rows_from_workbook():
    """Expected invalid rows derived FROM THE FILE: non-blank rows whose
    department or make cell is missing/sentinel — no hardcoded row numbers."""
    workbook = load_workbook(REFERENCE_WORKBOOK, read_only=True, data_only=True)
    sheet = workbook.active
    values = list(sheet.values)
    workbook.close()
    header = [cell if cell is not None else "" for cell in values[0]]
    missing_department: set[int] = set()
    missing_brand: set[int] = set()
    for row_number, raw_row in enumerate(values[1:], 2):
        record = dict(zip(header, raw_row))
        cells = list(record.values())
        if all(c is None or (isinstance(c, str) and not c.strip()) for c in cells):
            continue  # blank padding row — not a record
        if acs_mod._empty_to_none(record.get("department_name")) is None:
            missing_department.add(row_number)
        if acs_mod._empty_to_none(record.get("Brand / Company")) is None:
            missing_brand.add(row_number)
    return missing_department, missing_brand


@pytest.mark.skipif(not REFERENCE_WORKBOOK.exists(), reason="ACMS 301.xlsx reference fixture not present")
def test_reference_workbook_dry_run_matches_expected_classification(client):
    body = REFERENCE_WORKBOOK.read_bytes()
    data = _report(_post_dry(client, body, REFERENCE_WORKBOOK.name))

    # Physical vs actual record accounting (Excel max_row is not record count).
    assert data["physical_rows"] == 1021
    assert data["blank_rows_ignored"] == 721
    assert data["total_rows"] == 300
    assert data["invalid"] == 2
    assert data["imported"] == 298
    assert data["imported"] + data["invalid"] == data["total_rows"]

    # Header mapping: the make column ("Brand / Company") is recognized and
    # nothing in the file is left unmapped.
    assert "ac_make_company" in data["recognized_columns"]
    assert data["unknown_columns"] == []
    assert "ac_make_company" not in data["missing_fields"]

    # The genuinely invalid rows are derived from the file itself.
    missing_department, missing_brand = _derive_invalid_rows_from_workbook()
    assert missing_department and missing_brand  # fixture really contains both
    messages = {entry["row"]: entry["message"] for entry in data["errors"]}
    assert set(messages) == missing_department | missing_brand
    for row_number in missing_department:
        assert "department_name is required" in messages[row_number]
    for row_number in missing_brand:
        assert "Brand/company is required" in messages[row_number]
